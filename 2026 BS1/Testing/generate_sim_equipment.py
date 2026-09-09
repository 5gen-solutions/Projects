#!/usr/bin/env python3
"""
Generate sim_Equipment.XST + sim_Equipment.ST for BSR130 / BSR132.

Modelled on Inputs/sim_Equipment.XST (PLC703) patterns, using each site's
sim IO from the IO lists and naming conventions from generate_xst.py.
Includes ABB REX640 (BSR132) and FlexiSoft/PNOZ safety (FlexiSoft = BSR132 only).
"""
from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

from generate_xst import (
    CONTENT_DT,
    INPUTS,
    OUTPUTS,
    VERSION,
    Channel,
    Module,
    SiteData,
    load_site,
    now_dt,
    sim_tag,
    state_to_val,
    tag_base,
    var_ebool,
    var_typed,
    wrap_xst,
    xml_esc,
)

import openpyxl

PRODUCT = "Control Expert V16.2 - 250430"
SECTION_ORDER = 6


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def parse_dev(dev) -> tuple[str, Optional[int], str]:
    d = str(dev or "").strip().upper()
    m = re.match(r"([A-Z]+)(\d+)([A-Z]*)", d)
    if not m:
        return d, None, ""
    return m.group(1), int(m.group(2)), m.group(3)


def strip_dev_suffix(base: str, dev: str) -> str:
    """Equipment root = base with trailing device token removed."""
    b = base or ""
    d = str(dev or "").strip()
    if d and b.upper().endswith(d.upper()):
        return b[: -len(d)]
    # fallback: strip trailing alpha+digits
    return re.sub(r"[A-Za-z]+\d+[A-Za-z]*$", "", b)


def eng_defaults(base: str, desc) -> tuple[float, float]:
    """Heuristic engineering min/max for AHI_Scale."""
    u = (base or "").upper() + " " + str(desc or "").upper()
    if re.search(r"\b(TT|TEMP)", u):
        return 0.0, 120.0
    if re.search(r"\b(PT|PIT|PDT|PRESS)", u):
        return 0.0, 250.0
    if re.search(r"\b(FT|FIT|FLOW)", u):
        return 0.0, 100.0
    if re.search(r"\b(LT|LIT|LEVEL)", u):
        return 0.0, 100.0
    if re.search(r"\b(ZT|POS)", u):
        return 0.0, 100.0
    if re.search(r"\b(IT|CURRENT)", u):
        return 0.0, 1000.0
    if re.search(r"\b(XT|VOLT)", u):
        return 0.0, 1000.0
    if re.search(r"\b(WT|WEIGHT)", u):
        return 0.0, 5000.0
    return 0.0, 100.0


def do_tag(base: str) -> str:
    """Application DO / coil tag (non-sim), matching reference *_O style."""
    return f"{base}_O"


@dataclass
class AiChan:
    mod: Module
    ch: int
    chan: Channel
    pv_tag: str  # sim_<base>_PV
    ch_pv: str  # Sim_<mod>_Ch##PV


@dataclass
class CbPair:
    root: str
    cb_base: str
    trp_base: str
    label: str


@dataclass
class BrakeUnit:
    root: str
    release_sv: Optional[str]
    pump_sv: Optional[str]
    ps_lift: Optional[str]
    ps_lo: Optional[str]
    zs_lifts: list[str] = field(default_factory=list)


@dataclass
class RexChan:
    slot: str  # B / C / G
    iot: str
    ch: str  # BI1, PODP1, RF, SO1, ...
    device: str
    drive: str
    node: str
    desc: str
    state: str


@dataclass
class RexRelay:
    drive: str
    node: str  # PR1
    owner: str  # Drive+Node e.g. BMC1321INC01PR1
    label: str
    channels: list = field(default_factory=list)
    unmapped: list = field(default_factory=list)  # (slot, ch, reason)

    @property
    def prefix(self) -> str:
        return f"sim_{self.owner}"

    @property
    def is_incomer(self) -> bool:
        return "INC" in (self.drive or "").upper()


@dataclass
class FlexiSoftUnit:
    owner: str  # e.g. BCV131SLC1
    drive: str
    pri_ok_bases: list = field(default_factory=list)  # tag bases for PriOK pack
    notes: list = field(default_factory=list)


@dataclass
class PnozDevice:
    ntype: str
    instrument: str
    device: str
    drive: str
    desc: str
    sheet: str


@dataclass
class SafetyFeedback:
    base: str
    sim_i: str
    drive: str
    device: str
    desc: str
    sheet: str


@dataclass
class TeSysTUnit:
    """One Schneider TeSysT motor-starter / protection relay (Node Type TeSysT)."""
    drive: str
    owner: str  # Drive + PR1 (Rack=PR / Slot 1 → device DDT)
    label: str
    # PR-rack spreadsheet DI → sim_{Drive}{Device}_I (empty Instrument Tag → Drive+Device)
    cb_base: Optional[str] = None
    cax_base: Optional[str] = None
    sy_base: Optional[str] = None
    vac_base: Optional[str] = None
    cb_init: int = 1
    cax_init: int = 1
    sy_init: int = 1
    vac_init: int = 1

    @property
    def has_cb(self) -> bool:
        return bool(self.cb_base)

    @property
    def has_cax(self) -> bool:
        return bool(self.cax_base)

    @property
    def has_sy(self) -> bool:
        return bool(self.sy_base)

    @property
    def has_vac(self) -> bool:
        return bool(self.vac_base)


@dataclass
class GenStats:
    categories: dict = field(default_factory=dict)
    skipped: list = field(default_factory=list)
    rex_report: dict = field(default_factory=dict)
    safety_report: dict = field(default_factory=dict)
    tesyst_report: dict = field(default_factory=dict)
    init_ied_report: dict = field(default_factory=dict)


# ---------------------------------------------------------------------------
# Discovery
# ---------------------------------------------------------------------------

def collect_ai(site: SiteData) -> list[AiChan]:
    out: list[AiChan] = []
    for mod in site.modules:
        if mod.kind not in ("AHI", "ART"):
            continue
        for ch in range(8):
            c = mod.channels.get(ch)
            if not c or not c.base:
                continue
            if mod.kind == "AHI" and c.iot != "AI":
                continue
            if mod.kind == "ART" and c.iot not in ("RTDI", "AI"):
                # ART may be tagged RTDI
                if c.iot not in ("RTDI",):
                    continue
            out.append(
                AiChan(
                    mod=mod,
                    ch=ch,
                    chan=c,
                    pv_tag=sim_tag(c.base, "_PV"),
                    ch_pv=f"Sim_{mod.name}_Ch{ch:02d}PV",
                )
            )
    return out


def collect_cb_pairs(site: SiteData) -> list[CbPair]:
    cbs: dict[str, list[tuple[str, int, Channel]]] = defaultdict(list)
    trps: dict[str, list[tuple[str, int, Channel]]] = defaultdict(list)
    for c in site.all_di:
        if not c.base:
            continue
        code, num, _ = parse_dev(c.dev)
        if code == "CB" and num is not None:
            root = strip_dev_suffix(c.base, str(c.dev))
            cbs[root].append((c.base, num, c))
        elif code == "TRP" and num is not None:
            root = strip_dev_suffix(c.base, str(c.dev))
            trps[root].append((c.base, num, c))

    pairs: list[CbPair] = []
    used_cbs: set[str] = set()
    for root in sorted(set(cbs) & set(trps)):
        cb_list = sorted(cbs[root], key=lambda x: x[1])
        trp_list = sorted(trps[root], key=lambda x: x[1])
        for trp_base, trp_num, _ in trp_list:
            # prefer same-number CB
            cand = [x for x in cb_list if x[1] == trp_num and x[0] not in used_cbs]
            if not cand:
                cand = [x for x in cb_list if x[1] == 1 and x[0] not in used_cbs]
            if not cand:
                cand = [x for x in cb_list if x[0] not in used_cbs]
            if not cand:
                continue
            cb_base, cb_num, _ = cand[0]
            used_cbs.add(cb_base)
            pairs.append(
                CbPair(
                    root=root,
                    cb_base=cb_base,
                    trp_base=trp_base,
                    label=root or cb_base,
                )
            )
    return pairs


def collect_dos(site: SiteData) -> list[Channel]:
    out = []
    for m in site.modules:
        if m.kind != "DDO":
            continue
        for c in m.channels.values():
            if c.iot == "DO" and c.base:
                out.append(c)
    return out


def collect_brakes(site: SiteData, dos: list[Channel]) -> list[BrakeUnit]:
    """Find BKx equipment with SV + PS for SIM_Brake."""
    # group DI by BK root
    bk_di: dict[str, list[Channel]] = defaultdict(list)
    for c in site.all_di:
        b = c.base or ""
        m = re.search(r"(.+BK\d+)", b, re.I)
        if m:
            bk_di[m.group(1).upper()].append(c)

    bk_do: dict[str, list[Channel]] = defaultdict(list)
    for c in dos:
        b = c.base or ""
        m = re.search(r"(.+BK\d+)", b, re.I)
        if m:
            bk_do[m.group(1).upper()].append(c)

    units: list[BrakeUnit] = []
    for root in sorted(set(bk_di) | set(bk_do)):
        dis = bk_di.get(root, [])
        dos_r = bk_do.get(root, [])
        svs = [c for c in dos_r if parse_dev(c.dev)[0] == "SV"]
        svs = sorted(svs, key=lambda c: parse_dev(c.dev)[1] or 0)
        ps = [c for c in dis if parse_dev(c.dev)[0] == "PS"]
        # lift: state Lifted or desc Brake without Low
        ps_lift = None
        ps_lo = None
        for c in ps:
            st = str(c.state or "").casefold()
            desc = str(c.desc2 or "").casefold()
            # Prefer State over Desc2 (desc often says "Pressure Switch" for both)
            if st in {"low", "pressure low"} or (not st and "low" in desc):
                if ps_lo is None:
                    ps_lo = c.base
            elif st in {"lifted", "lift", "ok", "healthy", "high"} or "lift" in st:
                if ps_lift is None:
                    ps_lift = c.base
            elif "low" in desc and ps_lo is None:
                ps_lo = c.base
            elif ps_lift is None:
                ps_lift = c.base
        # If only one PS, treat as lift
        if ps_lift is None and ps and ps_lo is None:
            ps_lift = ps[0].base
        if ps_lift is None and len(ps) >= 1:
            # second pass: non-lo
            for c in ps:
                if c.base != ps_lo:
                    ps_lift = c.base
                    break
        # ZS lifted (A suffix often)
        zs_lifts = []
        for c in dis:
            code, num, suf = parse_dev(c.dev)
            if code == "ZS" and str(c.state or "").casefold() == "lifted":
                zs_lifts.append(c.base)
        if not svs or not (ps_lift or ps_lo):
            continue
        units.append(
            BrakeUnit(
                root=root,
                release_sv=svs[0].base if svs else None,
                pump_sv=svs[1].base if len(svs) > 1 else (svs[0].base if svs else None),
                ps_lift=ps_lift,
                ps_lo=ps_lo,
                zs_lifts=zs_lifts[:4],
            )
        )
    return units


def collect_gate_actuators(site: SiteData, dos: list[Channel]) -> list[dict]:
    """GA / FV with open+closed ZS and open/close SV."""
    results = []
    # Group ZS by root
    zs_by_root: dict[str, list[Channel]] = defaultdict(list)
    for c in site.all_di:
        if not c.base:
            continue
        code, _, _ = parse_dev(c.dev)
        if code != "ZS":
            continue
        root = strip_dev_suffix(c.base, str(c.dev))
        zs_by_root[root].append(c)

    sv_by_root: dict[str, list[Channel]] = defaultdict(list)
    for c in dos:
        code, _, _ = parse_dev(c.dev)
        if code != "SV":
            continue
        # For GA01SV01 base may be BAF130GA01SV01 — root BAF130GA01
        root = strip_dev_suffix(c.base, str(c.dev))
        if not root:
            root = str(c.drive or "")
        sv_by_root[root].append(c)

    for root, zslist in sorted(zs_by_root.items()):
        opens = [c for c in zslist if str(c.state or "").casefold() == "open"]
        closed = [c for c in zslist if str(c.state or "").casefold() == "closed"]
        if not opens or not closed:
            continue
        # Prefer roots that look like GA / FV / valve actuators
        if not re.search(r"(GA\d+|FV\d+)", root, re.I):
            # still allow if we have matching SVs
            if root not in sv_by_root:
                continue
        svs = sorted(sv_by_root.get(root, []), key=lambda c: parse_dev(c.dev)[1] or 0)
        if len(svs) < 1:
            continue
        results.append(
            {
                "root": root,
                "sv_open": svs[0].base,
                "sv_close": svs[1].base if len(svs) > 1 else None,
                "zs_open": [c.base for c in opens],
                "zs_closed": [c.base for c in closed],
            }
        )
    return results


def collect_cax_sr(site: SiteData, dos: list[Channel]) -> list[tuple[str, str]]:
    """Pairs (sim_CAX_I, SR_O) where same equipment root has CAX DI and SR DO."""
    cax: dict[str, str] = {}
    for c in site.all_di:
        if not c.base:
            continue
        code, _, _ = parse_dev(c.dev)
        if code == "CAX":
            root = strip_dev_suffix(c.base, str(c.dev))
            cax[root] = c.base
    sr: dict[str, str] = {}
    for c in dos:
        code, _, _ = parse_dev(c.dev)
        if code == "SR":
            root = strip_dev_suffix(c.base, str(c.dev))
            # also try drive
            sr[root] = c.base
            if c.drive:
                sr[str(c.drive).strip()] = c.base
    pairs = []
    for root, cax_b in sorted(cax.items()):
        if root in sr:
            pairs.append((cax_b, sr[root]))
    return pairs


def spray_groups(dos: list[Channel], ais: list[AiChan]) -> list[dict]:
    """Group water-addition / dust-suppression SVs and optional FT AI."""
    sprays = []
    for c in dos:
        desc = str(c.desc2 or "").casefold()
        code, _, _ = parse_dev(c.dev)
        if code != "SV":
            continue
        if any(
            k in desc
            for k in (
                "water addition",
                "dust suppression",
                "ore conditioning",
                "spray",
                "bulk ore",
            )
        ):
            sprays.append(c)
    if not sprays:
        return []

    # Group by drive / instrument family prefix
    by_key: dict[str, list[Channel]] = defaultdict(list)
    for c in sprays:
        drive = str(c.drive or "").strip()
        if drive:
            key = drive
        else:
            # BBN130SV01 -> BBN130
            key = re.sub(r"SV\d+[A-Z]*$", "", c.base or "", flags=re.I)
        by_key[key].append(c)

    groups = []
    for key, svs in sorted(by_key.items()):
        # Find FT AI sharing prefix
        ft = None
        for a in ais:
            bu = (a.chan.base or "").upper()
            if key.upper() in bu or bu.startswith(key[:6].upper()):
                if re.search(r"(FT|FIT)", bu):
                    ft = a
                    break
        # fallback: any FT on same first-3-letter family
        if ft is None:
            pref = key[:3].upper()
            for a in ais:
                if (a.chan.base or "").upper().startswith(pref) and re.search(
                    r"(FT|FIT)", (a.chan.base or "").upper()
                ):
                    ft = a
                    break
        groups.append({"key": key, "svs": sorted(svs, key=lambda x: x.base or ""), "ft": ft})
    return groups


# ---------------------------------------------------------------------------
# REX640 / FlexiSoft / PNOZ discovery from Sub + Field IO Assignment
# ---------------------------------------------------------------------------

def _sheet_cols(headers) -> dict[str, int]:
    cols: dict[str, int] = {}
    for i, h in enumerate(headers):
        if h and str(h).strip() and str(h).strip() not in cols:
            cols[str(h).strip()] = i
    return cols


def rex_sim_base(drive: str, device: str, instrument: str = "") -> Optional[str]:
    """Tag base for REX-local signals: prefer tag_base, else Drive+Device."""
    b = tag_base(instrument, device, drive)
    if b:
        return b
    drive = (drive or "").strip()
    device = (device or "").strip()
    if drive and device:
        return drive + device
    return None


def collect_rex640(site_name: str) -> list[RexRelay]:
    """Parse Sub IO Assignment for ABB REX640 relays (Node PR\\d+ / Node Type REX640)."""
    path = INPUTS / f"{site_name}_IO_List.xlsx"
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if "Sub IO Assignment" not in wb.sheetnames:
        wb.close()
        return []
    ws = wb["Sub IO Assignment"]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []
    cols = _sheet_cols(rows[0])

    def cell(r, name, default=""):
        i = cols.get(name)
        if i is None:
            return default
        v = r[i] if i < len(r) else None
        return str(v).strip() if v is not None else default

    relays: dict[tuple[str, str], RexRelay] = {}
    order: list[tuple[str, str]] = []
    current_key: Optional[tuple[str, str]] = None
    current_slot: Optional[str] = None

    for r in rows[1:]:
        node = cell(r, "Node")
        ntype = cell(r, "Node Type")
        drive = cell(r, "Drive")
        iot = cell(r, "IO Type")
        ch = cell(r, "Ch")
        dev = cell(r, "Device")
        desc = cell(r, "Desc2")
        state = cell(r, "State")
        it = cell(r, "Instrument Tag")
        is_pr = bool(re.fullmatch(r"PR\d+", node, re.I))
        is_rex_nt = "REX640" in ntype.upper()

        if is_pr and drive:
            key = (drive, node.upper())
            if key not in relays:
                owner = f"{drive}{node.upper()}"
                relays[key] = RexRelay(
                    drive=drive,
                    node=node.upper(),
                    owner=owner,
                    label=f"{drive} / {node.upper()}",
                    channels=[],
                )
                order.append(key)
                current_slot = None  # only clear slot when starting a NEW relay
            elif current_key != key:
                # switched to a different PR node
                current_slot = None
            current_key = key
        elif current_key and node and not re.fullmatch(r"PR\d+", node, re.I) and not is_rex_nt:
            # leaving PR block when Node jumps to something else
            current_key = None
            current_slot = None
            continue
        elif not current_key:
            continue

        if not current_key or current_key not in relays:
            continue

        m = re.search(r"SLOT\s*([BCG])", ntype, re.I)
        if m:
            current_slot = m.group(1).upper()

        if not iot:
            continue

        slot = current_slot
        chu = ch.upper()
        if chu.startswith("PODP") or chu == "RF":
            slot = "G"
        if slot is None:
            # cannot place
            relays[current_key].unmapped.append(("", ch or "?", "no slot context"))
            continue

        relays[current_key].channels.append(
            RexChan(
                slot=slot,
                iot=iot.upper(),
                ch=ch,
                device=dev,
                drive=relays[current_key].drive,
                node=relays[current_key].node,
                desc=desc,
                state=state,
            )
        )

    # Prefer only relays that look like REX (have SLOT markers or BI/PODP pattern)
    out: list[RexRelay] = []
    for key in order:
        rel = relays[key]
        has_bi = any(c.ch.upper().startswith("BI") for c in rel.channels)
        has_slot = any(c.slot in ("B", "C", "G") for c in rel.channels)
        if has_bi and has_slot:
            out.append(rel)
    return out


def _bit_to_word(dest: str, bit_exprs: dict[int, str]) -> str:
    parts = []
    for i in range(16):
        parts.append(f"BIT{i} := {bit_exprs.get(i, 'FALSE')}")
    inner = ",\n\t                         ".join(parts)
    return f"{dest} := \tBIT_TO_WORD ({inner});\n"


def pack_rex_slot_di(rel: RexRelay, slot: str) -> tuple[dict[int, str], list[tuple]]:
    """Map Slot B/C BI\\d+ → BIT(n-1); Slot G PODP1/2 → BIT7/8 (reference pattern)."""
    bits: dict[int, str] = {}
    unmapped = []
    for c in rel.channels:
        if c.slot != slot:
            continue
        if c.iot not in ("DI", "DIO"):
            continue
        chu = c.ch.upper()
        expr = None
        bit = None
        if slot in ("B", "C"):
            m = re.fullmatch(r"BI(\d+)", chu)
            if not m:
                unmapped.append((slot, c.ch, "non-BI DI"))
                continue
            bit = int(m.group(1)) - 1
            if bit < 0 or bit > 15:
                unmapped.append((slot, c.ch, "BI out of range"))
                continue
            if c.device:
                base = rex_sim_base(c.drive, c.device, "")
                expr = sim_tag(base, "_I") if base else "FALSE"
            else:
                expr = "FALSE"
                unmapped.append((slot, c.ch, "spare/empty device → FALSE"))
        elif slot == "G":
            if chu == "PODP1":
                bit = 7
            elif chu == "PODP2":
                bit = 8
            elif chu == "PODP3":
                unmapped.append((slot, c.ch, "PODP3 spare → FALSE"))
                continue
            else:
                unmapped.append((slot, c.ch, "unmapped G DI channel"))
                continue
            if c.device:
                base = rex_sim_base(c.drive, c.device, "")
                expr = sim_tag(base, "_I") if base else "FALSE"
            else:
                expr = "FALSE"
                unmapped.append((slot, c.ch, "empty device → FALSE"))
        if bit is not None and expr is not None:
            bits[bit] = expr
    return bits, unmapped


def collect_flexisoft_and_safety(site_name: str) -> tuple[list[FlexiSoftUnit], list[PnozDevice], list[SafetyFeedback]]:
    """Discover FlexiSoft Bus / PNOZ Node Types and safety-related DI feedbacks.

    FlexiSoft EIP simulation is owned by BSR132 only (Tim). When site_name is BSR132,
    also scan BSR130 IO lists for FlexiSoft/PWS primary elements that physically sit in
    BSR130 panels but communicate to the BSR132 PLC.
    """
    paths = [INPUTS / f"{site_name}_IO_List.xlsx"]
    if site_name == "BSR132":
        alt = INPUTS / "BSR130_IO_List.xlsx"
        if alt.exists():
            paths.append(alt)

    flexi_notes: dict[str, list[str]] = defaultdict(list)
    flexi_drives: set[str] = set()
    slc_ids: set[int] = set()
    pnoz: list[PnozDevice] = []
    feedbacks: list[SafetyFeedback] = []
    seen_fb: set[str] = set()
    pri_by_drive: dict[str, list[str]] = defaultdict(list)

    for path in paths:
        from_site = path.name.split("_")[0]
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        for sheet in ("Sub IO Assignment", "Field IO Assignment"):
            if sheet not in wb.sheetnames:
                continue
            rows = list(wb[sheet].iter_rows(values_only=True))
            if not rows:
                continue
            cols = _sheet_cols(rows[0])

            def cell(r, name, default=""):
                i = cols.get(name)
                if i is None:
                    return default
                v = r[i] if i < len(r) else None
                return str(v).strip() if v is not None else default

            for r in rows[1:]:
                ntype = cell(r, "Node Type")
                iot = cell(r, "IO Type")
                it = cell(r, "Instrument Tag")
                dev = cell(r, "Device")
                drive = cell(r, "Drive")
                desc = cell(r, "Desc2")
                ch = cell(r, "Ch")
                node = cell(r, "Node")
                blob = " ".join([ntype, iot, it, dev, drive, desc, ch, node]).upper()

                # FlexiSoft markers
                if (
                    "FLEXISOFT" in blob
                    or "FLEXI SOFT" in blob
                    or (dev.upper().startswith("SIL") and "FLEXI" in desc.upper())
                    or re.search(r"CPU\s*SLC", ch.upper())
                    or (str(ntype) == "1059305" and "SLC" in (ch + node).upper())
                ):
                    if drive:
                        flexi_drives.add(drive)
                        flexi_notes[drive].append(f"{from_site}/{sheet}: {dev or ntype} {desc}".strip())
                    slc_m = re.search(r"SLC\s*0*(\d+)", (ch + " " + node + " " + dev).upper())
                    if slc_m:
                        slc_ids.add(int(slc_m.group(1)))

                if "PNOZ" in ntype.upper():
                    # PNOZ inventory: only count rows from the site's own workbook
                    if from_site == site_name:
                        pnoz.append(
                            PnozDevice(
                                ntype=ntype,
                                instrument=it,
                                device=dev,
                                drive=drive,
                                desc=desc,
                                sheet=sheet,
                            )
                        )
                    if "FLEXI" in desc.upper() and drive:
                        flexi_drives.add(drive)
                        flexi_notes[drive].append(f"PNOZ→FlexiSoft ({from_site}): {it or drive} {desc}")

                # Safety DI feedbacks — site's own workbook only (tags belong to that PLC)
                if from_site == site_name and iot.upper() == "DI" and (it or dev):
                    du = (dev or "").upper()
                    desc_u = (desc or "").upper()
                    is_safety = (
                        du.startswith(("SY", "PWS", "LOS", "LCS"))
                        or "SAFETY" in desc_u
                        or "PULLWIRE" in desc_u
                        or "EMERGENCY" in desc_u
                        or "SLC" in du
                    )
                    if is_safety:
                        base = tag_base(it, dev, drive)
                        if not base and drive and dev:
                            base = drive + dev
                        if not base and it:
                            base = it
                        if base and base not in seen_fb:
                            seen_fb.add(base)
                            feedbacks.append(
                                SafetyFeedback(
                                    base=base,
                                    sim_i=sim_tag(base, "_I"),
                                    drive=drive or "",
                                    device=dev or "",
                                    desc=desc or "",
                                    sheet=sheet,
                                )
                            )
                            if re.match(r"(PWS|LOS)\d*", du) or re.search(r"(PWS|LOS)\d+", (it or "").upper()):
                                pri_by_drive[drive or ""].append(base)
                            # SLC01 / SLC02 markers on device names
                            sm = re.search(r"SLC\s*0*(\d+)", du) or re.search(r"SLC\s*0*(\d+)", base.upper())
                            if sm:
                                slc_ids.add(int(sm.group(1)))
                                if drive:
                                    flexi_drives.add(drive)

                # Primary elements from cross-site (BSR130) for BSR132 FlexiSoft PriOK only
                if site_name == "BSR132" and from_site != site_name and iot.upper() == "DI" and (it or dev):
                    du = (dev or "").upper()
                    if re.match(r"(PWS|LOS)\d*", du) or re.search(r"(PWS|LOS)\d+", (it or "").upper()):
                        base = tag_base(it, dev, drive)
                        if not base and drive and dev:
                            base = drive + dev
                        if base:
                            pri_by_drive[drive or ""].append(base)
                            # note only — do not add to this site's feedback/init list
        wb.close()

    # Build FlexiSoft units (caller clears these for non-BSR132)
    flexi: list[FlexiSoftUnit] = []
    if flexi_drives or slc_ids:
        drive = sorted(flexi_drives)[0] if flexi_drives else "BCV131"
        # Prefer BCV131 if present
        for d in flexi_drives:
            if d.upper().startswith("BCV131"):
                drive = d
                break
        ids = sorted(slc_ids) if slc_ids else [1]
        # Always at least SLC1 when FlexiSoft bus present
        if 1 not in ids:
            ids = [1] + ids
        for sid in ids:
            owner = f"{drive}SLC{sid}"
            unit = FlexiSoftUnit(owner=owner, drive=drive, notes=list(flexi_notes.get(drive, []))[:8])
            unit.notes.append(f"SLC{sid:02d} / owner {owner} (FlexiSoft → BSR132 PLC)")
            # PriOK: PWS/LOS from related drives
            drives_wanted = {drive}
            for d in list(pri_by_drive.keys()):
                du = d.upper()
                if du.startswith("BCV") or du.startswith("BAF") or du.startswith("BBN"):
                    drives_wanted.add(d)
            bases: list[str] = []
            seen: set[str] = set()
            for d in sorted(drives_wanted):
                for b in sorted(set(pri_by_drive.get(d, []))):
                    if b not in seen:
                        seen.add(b)
                        bases.append(b)
            if len(bases) < 4:
                for fb in feedbacks:
                    if fb.drive in drives_wanted and re.search(r"(LCS|PWS|LOS)", fb.base.upper()):
                        if fb.base not in seen:
                            seen.add(fb.base)
                            bases.append(fb.base)
            unit.pri_ok_bases = bases
            flexi.append(unit)

    return flexi, pnoz, feedbacks



def _tesyst_dev_base(drive: str, device: str, instrument: str = "") -> Optional[str]:
    """Tag base for TeSysT / PR-rack DI: tag_base rules, else Drive+Device."""
    b = tag_base(instrument, device, drive)
    if b:
        return b
    drive = (drive or "").strip()
    device = (device or "").strip()
    if drive and device:
        return drive + device
    return None


def collect_tesyst(site_name: str) -> list[TeSysTUnit]:
    """Parse Sub IO Assignment for Node Type TeSysT (one unit per Drive).

    IO lists show Rack=PR with CB01 (Node Type TeSysT) plus sibling CAX01/SY01/(VAC01)
    DI on the same Drive/PR rack. Device DDT owner: {Drive}PR1.
    Sim tags use Drive+Device when Instrument Tag is empty (excluded from X80 IOmap).
    """
    path = INPUTS / f"{site_name}_IO_List.xlsx"
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if "Sub IO Assignment" not in wb.sheetnames:
        wb.close()
        return []
    ws = wb["Sub IO Assignment"]
    rows = list(ws.iter_rows(values_only=True))
    cols = _sheet_cols(rows[0]) if rows else {}

    def cell(r, name, default=""):
        i = cols.get(name)
        if i is None:
            return default
        v = r[i] if i < len(r) else None
        return str(v).strip() if v is not None else default

    drives: dict[str, TeSysTUnit] = {}
    order: list[str] = []
    for r in rows[1:]:
        ntype = cell(r, "Node Type")
        node = cell(r, "Node")
        if "TESYST" not in ntype.upper() and node.upper() != "TESYST":
            continue
        drive = cell(r, "Drive")
        if not drive:
            continue
        if drive not in drives:
            rack = cell(r, "Rack").upper()
            slot = cell(r, "Slot")
            if rack == "PR" or not rack:
                owner = f"{drive}PR1"
            else:
                owner = f"{drive}{rack}{slot or '1'}"
            drives[drive] = TeSysTUnit(
                drive=drive,
                owner=owner,
                label=f"{drive} TeSysT → {owner}",
            )
            order.append(drive)

    # PR-rack DI on TeSysT drives: CB01 / CAX01 / SY01 / VAC01
    for r in rows[1:]:
        drive = cell(r, "Drive")
        if drive not in drives:
            continue
        if cell(r, "Rack").upper() != "PR":
            continue
        if cell(r, "IO Type").upper() != "DI":
            continue
        dev = cell(r, "Device")
        if not dev:
            continue
        code, _, _ = parse_dev(dev)
        if code not in ("CB", "CAX", "SY", "VAC"):
            continue
        # Prefer primary devices CB01/CAX01/SY01/VAC01 (ignore CAX02 etc. for force pins)
        if code == "CAX" and not re.fullmatch(r"CAX01", dev, re.I):
            continue
        if code == "CB" and not re.fullmatch(r"CB01", dev, re.I):
            continue
        if code == "SY" and not re.fullmatch(r"SY01", dev, re.I):
            continue
        if code == "VAC" and not re.fullmatch(r"VAC01", dev, re.I):
            continue
        it = cell(r, "Instrument Tag")
        state = cell(r, "State")
        base = _tesyst_dev_base(drive, dev, it)
        if not base:
            continue
        u = drives[drive]
        init_v = state_to_val(state, dev, None)
        if code == "CB" and not u.cb_base:
            u.cb_base = base
            u.cb_init = init_v
        elif code == "CAX" and not u.cax_base:
            u.cax_base = base
            u.cax_init = init_v
        elif code == "SY" and not u.sy_base:
            u.sy_base = base
            u.sy_init = init_v
        elif code == "VAC" and not u.vac_base:
            u.vac_base = base
            u.vac_init = init_v

    wb.close()
    return [drives[d] for d in order]


def emit_tesyst_st(units: list[TeSysTUnit], need, stats: GenStats) -> list[str]:
    """Emit SIM_TeSysT instances — DDT in/out like Simocode; force pins from sim_ IO."""
    st: list[str] = []
    st.append(banner("TeSysT motor starters (SIM_TeSysT)"))
    if not units:
        st.append("(* No TeSysT Node Type rows in Sub IO Assignment *)\n")
        stats.categories["SIM_TeSysT"] = 0
        return st

    need("sim_TeSysT_MstIP", "IPAddDDT", "Shared TeSysT master IP (PLC NOC / EIP)")
    need("sim_init", "EBOOL", "One-shot initialisation flag (shared with sim_Initialise)")

    # Healthy defaults for TeSysT PR sim DI (also mirrored in sim_Initialise)
    st.append("(* TeSysT PR sim DI healthy defaults (CB Closed, SY/VAC Healthy, CAX Closed) *)")
    st.append("IF sim_init THEN")
    for u in units:
        if u.cb_base:
            tag = sim_tag(u.cb_base, "_I")
            need(tag, "EBOOL", f"{u.cb_base} TeSysT CB closed/healthy")
            st.append(f"\t{tag} := {u.cb_init};")
        if u.cax_base:
            tag = sim_tag(u.cax_base, "_I")
            need(tag, "EBOOL", f"{u.cax_base} TeSysT contactor aux")
            st.append(f"\t{tag} := {u.cax_init};")
        if u.sy_base:
            tag = sim_tag(u.sy_base, "_I")
            need(tag, "EBOOL", f"{u.sy_base} TeSysT safety interlock")
            st.append(f"\t{tag} := {u.sy_init};")
        if u.vac_base:
            tag = sim_tag(u.vac_base, "_I")
            need(tag, "EBOOL", f"{u.vac_base} TeSysT 240VAC supply")
            st.append(f"\t{tag} := {u.vac_init};")
    st.append("\t(* Master IP left 0.0.0.0 unless set by sim_Initialise / HMI *)")
    st.append("END_IF;\n")

    for i, u in enumerate(units):
        st.append(f"(* {u.label} *)")
        need(f"SIM_TeSysT_{i}", "SIM_TeSysT")

        if u.cb_base:
            cb = f"NOT {sim_tag(u.cb_base, '_I')}"
            need(sim_tag(u.cb_base, "_I"), "EBOOL", f"{u.cb_base} TeSysT CB closed/healthy")
        else:
            cb = "FALSE"  # channel absent → treat as closed/healthy (not open)

        if u.vac_base:
            vac = f"NOT {sim_tag(u.vac_base, '_I')}"
            need(sim_tag(u.vac_base, "_I"), "EBOOL", f"{u.vac_base} TeSysT 240VAC supply")
        else:
            vac = "FALSE"

        if u.cax_base:
            cax_tag = sim_tag(u.cax_base, "_I")
            need(cax_tag, "EBOOL", f"{u.cax_base} TeSysT contactor aux")
            cax = f"BOOL_TO_INT(IN := {cax_tag})"
        else:
            cax = "0"  # channel truly absent

        if u.sy_base:
            ess = f"NOT {sim_tag(u.sy_base, '_I')}"
            need(sim_tag(u.sy_base, "_I"), "EBOOL", f"{u.sy_base} TeSysT safety interlock")
        else:
            ess = "FALSE"

        st.append(
            f"SIM_TeSysT_{i} (iOutDatDDT := {u.owner}.Outputs,\n"
            f"               iMstIPAdr := sim_TeSysT_MstIP,\n"
            f"               iRun := 0,\t(* 0=Sim; 1=Stop; 2=Fwd; 3=Rev *)\n"
            f"               iCB01_Opn := {cb},\t(* Open pin ← NOT Closed sim DI *)\n"
            f"               iVAC01_Off := {vac},\n"
            f"               iCAX01 := {cax},\t(* INT from sim CAX DI *)\n"
            f"               iESS01_Off := {ess},\n"
            f"               iFltID := 0,\n"
            f"               oInDatDDT => {u.owner}.Inputs);\n"
        )

    stats.categories["SIM_TeSysT"] = len(units)
    stats.tesyst_report = {
        "count": len(units),
        "owners": [u.owner for u in units],
        "drives": [u.drive for u in units],
        "cb": sum(1 for u in units if u.has_cb),
        "cax": sum(1 for u in units if u.has_cax),
        "sy": sum(1 for u in units if u.has_sy),
        "vac": sum(1 for u in units if u.has_vac),
        "wiring": (
            "iOutDatDDT/oInDatDDT ↔ {Drive}PR1.Outputs/.Inputs; "
            "iMstIPAdr ← sim_TeSysT_MstIP; iRun=0 Sim; iFltID=0; "
            "iCB01_Opn := NOT sim_{Drive}CB01_I; "
            "iCAX01 := BOOL_TO_INT(IN := sim_{Drive}CAX01_I); "
            "iESS01_Off := NOT sim_{Drive}SY01_I; "
            "iVAC01_Off := NOT sim_{Drive}VAC01_I when present else FALSE."
        ),
    }
    return st


def emit_rex640_st(relays: list[RexRelay], need, stats: GenStats) -> list[str]:
    st: list[str] = []
    if not relays:
        stats.categories["sim_rex640"] = 0
        stats.rex_report = {"count": 0, "owners": [], "unmapped": []}
        return st  # no REX640 on this site — omit N/A commentary

    st.append(banner("ABB REX640 6.6kV Relays"))

    need("sim_init", "EBOOL", "One-shot initialisation flag (shared with sim_Initialise)")
    all_unmapped = []
    owners = []
    for i, rel in enumerate(relays):
        owners.append(rel.owner)
        label = "Incomer" if rel.is_incomer else f"Feeder ({rel.node})"
        st.append(f"(* {label}: {rel.label} → {rel.owner}.Inputs *)")
        # Slot B / C / G DI packs
        for slot, dest_suffix in (("B", "SlotBDI"), ("C", "SlotCDI"), ("G", "SlotGDIO")):
            bits, unm = pack_rex_slot_di(rel, slot)
            all_unmapped.extend([(rel.owner, s, ch, why) for s, ch, why in unm])
            dest = f"{rel.prefix}_{dest_suffix}"
            need(dest, "WORD", f"{rel.owner} REX640 Slot {slot} digital in pack")
            # declare DI tags used
            for expr in bits.values():
                if expr != "FALSE" and expr.startswith("sim_"):
                    need(expr, "EBOOL")
            st.append(_bit_to_word(dest, bits))

        # DO / status words + analogs
        for suffix, typ, cmt in (
            ("SlotBDO", "WORD", "Slot B digital out pack"),
            ("SlotCDO", "WORD", "Slot C digital out pack"),
            ("Flt", "EBOOL", "Forced fault"),
            ("FunSts", "WORD", "Function status"),
            ("Amps", "REAL", "Simulated current"),
            ("Volts", "REAL", "Simulated voltage kV"),
        ):
            need(f"{rel.prefix}_{suffix}", typ, f"{rel.owner} {cmt}")

        st.append(
            f"IF sim_init THEN\n"
            f"\t{rel.prefix}_Amps\t:= 15.0;\n"
            f"\t{rel.prefix}_Volts\t:= 6.6;\n"
            f"END_IF;"
        )
        i_in = "600.0" if rel.is_incomer else "400.0"
        need(f"sim_rex640_{i}", "sim_rex640")
        st.append(
            f"sim_rex640_{i} (iFlt \t:= {rel.prefix}_Flt,\n"
            f"              iSlotBDI \t:= {rel.prefix}_SlotBDI,\n"
            f"              iSlotCDI \t:= {rel.prefix}_SlotCDI,\n"
            f"              iSlotBDO := {rel.prefix}_SlotBDO,\n"
            f"              iSlotCDO := {rel.prefix}_SlotCDO,\n"
            f"              iSlotGDIO := {rel.prefix}_SlotGDIO,\n"
            f"              iFunSts := {rel.prefix}_FunSts,\n"
            f"              iAmps := {rel.prefix}_Amps,\n"
            f"              iVolts := {rel.prefix}_Volts,\n"
            f"\t      iIn\t:= {i_in},\n"
            f"\t      iUn\t:= 6.6,\n"
            f"              oDatDDT => {rel.owner}.Inputs);\n"
        )

    stats.categories["sim_rex640"] = len(relays)
    stats.rex_report = {
        "count": len(relays),
        "owners": owners,
        "unmapped": all_unmapped,
        "bit_mapping": (
            "Slot B/C: BI{n} → BIT{n-1} via sim_{Drive}{Device}_I; empty device → FALSE. "
            "Slot G: PODP1 → BIT7, PODP2 → BIT8 (reference PLC703 pattern); other G bits FALSE."
        ),
    }
    return st


def emit_safety_st(
    flexi: list[FlexiSoftUnit],
    pnoz: list[PnozDevice],
    feedbacks: list[SafetyFeedback],
    need,
    stats: GenStats,
) -> list[str]:
    st: list[str] = []
    if flexi:
        st.append(banner("Safety — FlexiSoft / PNOZ / interlock feedbacks"))
    else:
        st.append(banner("Safety — PNOZ / interlock feedbacks"))


    # Inventory comments
    if pnoz:
        st.append("(* PNOZ devices from Sub + Field IO Assignment: *)")
        by = defaultdict(list)
        for p in pnoz:
            by[p.ntype].append(p)
        for nt in sorted(by):
            items = by[nt]
            drives = sorted({(p.drive or p.instrument or "?") for p in items})
            st.append(f"(*   {nt}: {len(items)} row(s) — drives/tags: {', '.join(drives[:12])} *)")
        st.append("")

    # FlexiSoft SIM blocks
    if flexi:
        need("sim_init", "EBOOL", "One-shot initialisation flag (shared with sim_Initialise)")
        for i, unit in enumerate(flexi):
            st.append(f"(* FlexiSoft {unit.owner} (drive {unit.drive}) *)")
            for note in unit.notes[:6]:
                st.append(f"(*   {note} *)")
            arr = f"sim_{unit.owner}_PriOKArr"
            need(arr, "ARRAY[0..3] OF WORD", f"{unit.owner} primary-element healthy words")
            need(f"SIM_FlexiSoft_{i}", "SIM_FlexiSoft")

            # Pack PriOK: 4 words × 16 bits; unused → TRUE (reference pads TRUE)
            bases = list(unit.pri_ok_bases)
            for wi in range(4):
                bits = {}
                for bi in range(16):
                    idx = wi * 16 + bi
                    if idx < len(bases):
                        tag = sim_tag(bases[idx], "_I")
                        need(tag, "EBOOL")
                        bits[bi] = tag
                    else:
                        bits[bi] = "TRUE"
                parts = ",\n                         ".join(f"BIT{b} := {bits[b]}" for b in range(16))
                st.append(f"{arr}[{wi}] := BIT_TO_WORD\n\t\t\t({parts});")

            st.append(
                f"SIM_FlexiSoft_{i} (iOutDatDDT := {unit.owner}.Outputs,\n"
                f"                 iPriOKWd1 := {arr}[0],\n"
                f"                 iPriOKWd2 := {arr}[1],\n"
                f"\t\t iPriOKWd3 := {arr}[2],\n"
                f"                 iPriOKWd4 := {arr}[3],\n"
                f"                 oInDatDDT => {unit.owner}.Inputs);\n"
            )
        stats.categories["SIM_FlexiSoft"] = len(flexi)
    else:
        stats.categories["SIM_FlexiSoft"] = 0

    # Hold safety interlock / LCS / SY feedbacks healthy on sim_init (PNOZ & hardwired)
    safety_init = [
        fb
        for fb in feedbacks
        if re.search(r"(SY\d*|SLC\d*SY\d*)", fb.device.upper())
        or re.search(r"(LCS\d+.*SY|SY\d+)$", fb.base.upper())
        or "SAFETY INTERLOCK" in fb.desc.upper()
        or "SAFETY RELAY" in fb.desc.upper()
        or "EMERGENCY" in fb.desc.upper()
        or "SAFETY SLC" in fb.desc.upper()
        or "SAFETY PLC" in fb.desc.upper()
    ]
    # Always include PWS healthy bits as init defaults too when present
    pws_fb = [fb for fb in feedbacks if re.match(r"PWS\d+", (fb.device or "").upper()) or re.search(r"PWS\d+", fb.base.upper())]

    init_list = []
    seen = set()
    for fb in safety_init + pws_fb:
        if fb.sim_i not in seen:
            seen.add(fb.sim_i)
            init_list.append(fb)

    if init_list:
        st.append("(* Safety / PNOZ-related DI healthy defaults (also set in sim_Initialise) *)")
        st.append("IF sim_init THEN")
        need("sim_init", "EBOOL")
        for fb in init_list:
            need(fb.sim_i, "EBOOL", f"{fb.base} — {fb.desc}")
            st.append(f"\t{fb.sim_i} := 1;")
        st.append("END_IF;\n")
        stats.categories["safety_DI_init"] = len(init_list)
    else:
        st.append("(* No SY/LCS/PWS safety DI feedbacks found to initialise *)\n")
        stats.categories["safety_DI_init"] = 0

    stats.categories["PNOZ_rows"] = len(pnoz)
    stats.safety_report = {
        "flexi_owners": [u.owner for u in flexi],
        "flexi_pri_ok_counts": {u.owner: len(u.pri_ok_bases) for u in flexi},
        "flexi_pri_ok_bases": {u.owner: u.pri_ok_bases for u in flexi},
        "pnoz_count": len(pnoz),
        "pnoz_types": sorted({p.ntype for p in pnoz}),
        "pnoz_drives": sorted({p.drive or p.instrument for p in pnoz if (p.drive or p.instrument)}),
        "safety_feedback_count": len(feedbacks),
        "safety_init_count": len(init_list),
        "mapping": (
            "FlexiSoft: SIM_FlexiSoft with owner {Drive}SLC1.Outputs/.Inputs; "
            "PriOK words packed sequentially from PWS/LOS/(LCS) DI tag bases for related drives; "
            "unused bits TRUE. PNOZ: inventory from Node Type; associated SY/LCS/PWS DIs held healthy on sim_init."
        ),
    }
    return st


# ---------------------------------------------------------------------------
# ST emission
# ---------------------------------------------------------------------------

def banner(title: str) -> str:
    return (
        "(* ====================================\n"
        f"\n"
        f"     {title}\n"
        f"\n"
        f"=======================================*)\n"
    )


def emit_equipment(site: SiteData) -> tuple[str, str, GenStats]:
    """Return (st_source, data_block_xml, stats)."""
    stats = GenStats()
    rex_relays = collect_rex640(site.site)
    flexi_units, pnoz_devs, safety_fbs = collect_flexisoft_and_safety(site.site)

    # Tim: ALL FlexiSoft I/O communicates to BSR132 PLC only — never emit SIM_FlexiSoft on BSR130.
    if site.site != "BSR132":
        if flexi_units:
            stats.skipped.append(
                "SIM_FlexiSoft — FlexiSoft rows may appear in this site IO list / panels, "
                "but FlexiSoft communicates to BSR132 PLC only (per Tim); omitted here"
            )
        flexi_units = []

    # Keep skip notes for the console report only; do NOT dump site-N/A noise into BSR130 ST.
    if site.site == "BSR132":
        stats.skipped.extend([
            "SIM_SimocodeEIP3 — no Simocode EIP device DDTs; CB/TRP handled via sim_CBSts",
            "SIM_HarmonyNXGProEIP — no Harmony VSD EIP devices in hardwired IO",
            "sim_X80EHC0200 — EHC speed-sim already in sim_PLC_IOmap (sim_X80EHC0800)",
            "PROC_Linear — no clear pump/cooler Run feedback tags to drive rates; AHI_Scale used for AI/ART PVs instead",
        ])
        if not flexi_units:
            stats.skipped.insert(
                0, "SIM_FlexiSoft — no FlexiSoft Bus / SLC mapping found for BSR132"
            )
    # BSR130: no FlexiSoft/REX "not on this PLC" skip commentary (per Tim cleanup)

    ais = collect_ai(site)
    dos = collect_dos(site)
    cb_pairs = collect_cb_pairs(site)
    brakes = collect_brakes(site, dos)
    gates = collect_gate_actuators(site, dos)
    cax_sr = collect_cax_sr(site, dos)
    sprays = spray_groups(dos, ais)
    tesyst_units = collect_tesyst(site.site)

    # Variables to declare
    db_names: dict[str, tuple[str, Optional[str]]] = {}  # name -> (type, comment)

    def need(name: str, typ: str, comment: str | None = None):
        if name not in db_names:
            db_names[name] = (typ, comment)

    st: list[str] = []
    st.append(banner("Equipment Simulations - START"))
    st.append("")

    # Random
    st.append("(* Constant random flow *)")
    st.append(
        "Simulation_Random_Number_0 (\tLowerLimit \t:= -0.1,\n"
        "                            \tUpperLimit \t:= 0.1,\n"
        "                            \tRandomNumber \t=> Sim_FTxx_rand);\n"
    )
    need("Simulation_Random_Number_0", "SIM_Random_Number")
    need("Sim_FTxx_rand", "REAL", "Shared random noise for equipment sims")

    st.append("(* enable and disable all simulations *)")
    st.append("IF Sim_EquipBlocks THEN\n")
    need("Sim_EquipBlocks", "EBOOL", "Enable equipment simulation blocks")

    # ---- ABB REX640 (near top, matching reference) ----
    st.extend(emit_rex640_st(rex_relays, need, stats))

    # ---- FlexiSoft (BSR132 only) / PNOZ / safety ----
    st.extend(emit_safety_st(flexi_units, pnoz_devs, safety_fbs, need, stats))

    # ---- TeSysT motor starters ----
    st.extend(emit_tesyst_st(tesyst_units, need, stats))

    # Skipped stubs — BSR132 only (BSR130: no "not on this PLC" SKIP noise)
    if site.site == "BSR132" and stats.skipped:
        st.append(banner("Reference patterns not applicable (skipped)"))
        for s in stats.skipped:
            st.append(f"(* SKIP: {s} *)")
        st.append("")

    # ---- Analog AI scaling ----
    st.append(banner("Analog AI / ART scaling (sim_AHI_Scale)"))
    stats.categories["sim_AHI_Scale"] = len(ais)
    for i, a in enumerate(ais):
        emin, emax = eng_defaults(a.chan.base or "", a.chan.desc2)
        st.append(f"(* {a.chan.base} — {a.chan.desc2 or ''} → {a.ch_pv} *)")
        st.append(
            f"sim_AHI_Scale_{i} (ioPV := {a.pv_tag},\n"
            f"                 iMax := {emax},\n"
            f"                 iMin := {emin},\n"
            f"                 iEngMin := {emin},\n"
            f"                 iEngMax := {emax},\n"
            f"                 iRand_LL := -0.05,\n"
            f"                 iRand_UL := 0.05,\n"
            f"                 oPVPer => {a.ch_pv});\n"
        )
        need(f"sim_AHI_Scale_{i}", "sim_AHI_Scale")
        need(a.pv_tag, "REAL", f"{a.chan.base} engineering PV")
        need(a.ch_pv, "REAL", f"{a.mod.name} Ch{a.ch:02d} percent PV")
    if not ais:
        st.append("(* No AHI/ART channels with instrument tags *)\n")

    # ---- Water addition / sprays ----
    st.append(banner("Water addition / dust suppression / spray valves"))
    spray_sv_count = 0
    for g in sprays:
        key = g["key"]
        svs = g["svs"]
        ft: Optional[AiChan] = g["ft"]
        st.append(f"(*---- {key} sprays ({len(svs)} SV) ----*)")
        parts = []
        for sv in svs:
            spray_sv_count += 1
            pv = sim_tag(sv.base, "_PV")
            o = do_tag(sv.base)
            need(o, "EBOOL", f"{sv.base} solenoid output")
            need(pv, "REAL")
            st.append(
                f"IF {o} THEN\n"
                f"\t{pv} := 1.0;  (* nominal spray contribution *)\n"
                f"ELSE\n"
                f"\t{pv} := 0.0;\n"
                f"END_IF;\n"
            )
            parts.append(pv)
        if ft and parts:
            sum_expr = " + ".join(parts)
            # scale nominal: each SV = portion of eng max
            st.append(f"{ft.pv_tag} := ({sum_expr}) * 10.0;  (* L/s-ish aggregate into {ft.chan.base} *)\n")
            need(ft.pv_tag, "REAL", f"{ft.chan.base} engineering PV")
        elif parts:
            st.append(f"(* No FT AI counterpart for {key} — SV PVs only *)\n")
    stats.categories["spray_SV"] = spray_sv_count
    if not sprays:
        st.append("(* No spray / water-addition SV DOs found *)\n")

    # ---- Gate / valve actuators ----
    st.append(banner("Gate / valve actuator position feedback"))
    stats.categories["gate_valve_ST"] = len(gates)
    for g in gates:
        st.append(f"(* {g['root']} *)")
        o_open = do_tag(g["sv_open"])
        need(o_open, "EBOOL")
        zs_o = g["zs_open"]
        zs_c = g["zs_closed"]
        for z in zs_o + zs_c:
            need(sim_tag(z, "_I"), "EBOOL")
        if g["sv_close"]:
            o_close = do_tag(g["sv_close"])
            need(o_close, "EBOOL")
            open_assigns = "\n".join(f"\t{sim_tag(z, '_I')} := 1;" for z in zs_o)
            open_clears = "\n".join(f"\t{sim_tag(z, '_I')} := 0;" for z in zs_c)
            close_assigns = "\n".join(f"\t{sim_tag(z, '_I')} := 1;" for z in zs_c)
            close_clears = "\n".join(f"\t{sim_tag(z, '_I')} := 0;" for z in zs_o)
            st.append(
                f"IF {o_open} THEN\n"
                f"{open_assigns}\n"
                f"{open_clears}\n"
                f"ELSIF {o_close} THEN\n"
                f"{close_assigns}\n"
                f"{close_clears}\n"
                f"END_IF;\n"
            )
        else:
            # single SV: energised = open
            open_assigns = "\n".join(f"\t{sim_tag(z, '_I')} := 1;" for z in zs_o)
            open_clears = "\n".join(f"\t{sim_tag(z, '_I')} := 0;" for z in zs_c)
            close_assigns = "\n".join(f"\t{sim_tag(z, '_I')} := 1;" for z in zs_c)
            close_clears = "\n".join(f"\t{sim_tag(z, '_I')} := 0;" for z in zs_o)
            st.append(
                f"IF {o_open} THEN\n"
                f"{open_assigns}\n"
                f"{open_clears}\n"
                f"ELSE\n"
                f"{close_assigns}\n"
                f"{close_clears}\n"
                f"END_IF;\n"
            )
    if not gates:
        st.append("(* No GA/FV open-close SV+ZS pairs found on DDO *)\n")
    # Stub: FV with open/closed ZS but no local SV DO (AO/remote drive)
    fv_stubbed = []
    zs_by_root = {}
    from collections import defaultdict as _dd
    zs_by_root = _dd(list)
    for c in site.all_di:
        if not c.base:
            continue
        code, _, _ = parse_dev(c.dev)
        if code != "ZS":
            continue
        root = strip_dev_suffix(c.base, str(c.dev))
        zs_by_root[root].append(c)
    gated_roots = {g["root"] for g in gates}
    for root, zslist in sorted(zs_by_root.items()):
        if root in gated_roots:
            continue
        if not re.search(r"FV\d+", root, re.I):
            continue
        opens = [c.base for c in zslist if str(c.state or "").casefold() == "open"]
        closed = [c.base for c in zslist if str(c.state or "").casefold() == "closed"]
        if opens and closed:
            fv_stubbed.append(root)
            st.append(
                f"(* STUB: {root} has ZS open={opens} closed={closed} "
                f"but no SV DO on this PLC — drive from AO/remote or HMI *)"
            )
    if fv_stubbed:
        stats.categories["FV_stub_comment"] = len(fv_stubbed)
        st.append("")

    # ---- Brakes ----
    st.append(banner("Conveyor / tail brake (SIM_Brake)"))
    stats.categories["SIM_Brake"] = len(brakes)
    for i, b in enumerate(brakes):
        st.append(f"(* {b.root} *)")
        rel = do_tag(b.release_sv) if b.release_sv else "FALSE"
        pmp = do_tag(b.pump_sv) if b.pump_sv else rel
        if b.release_sv:
            need(do_tag(b.release_sv), "EBOOL")
        if b.pump_sv:
            need(do_tag(b.pump_sv), "EBOOL")
        flt = f"sim_{b.root}_FltMod"
        rst = f"sim_{b.root}_Rst"
        prs = f"sim_{b.root}_Prs_PV"
        need(flt, "INT", f"{b.root} brake fault mode")
        need(rst, "EBOOL", f"{b.root} brake reset")
        need(prs, "REAL", f"{b.root} brake pressure PV")
        need(f"SIM_Brake_{i}", "SIM_Brake")
        o_lft = sim_tag(b.ps_lift, "_I") if b.ps_lift else f"sim_{b.root}_Lft_I"
        o_lo = sim_tag(b.ps_lo, "_I") if b.ps_lo else f"sim_{b.root}_PrsLo_I"
        need(o_lft, "EBOOL")
        need(o_lo, "EBOOL")
        st.append(
            f"SIM_Brake_{i} (iReleaseSV \t:= {rel}\t,\n"
            f"             iPmpReq \t\t:= {pmp}\t,\n"
            f"             iBrakePrsSP \t:= 1000.0\t\t,\n"
            f"             iBrakePrsHoldSP \t:= 1050.0\t\t,\n"
            f"             iLiftTime \t\t:= t#5s\t\t\t,(* sim faster *)\n"
            f"             iDumpTime \t\t:= t#5s\t\t\t,\n"
            f"             iFltMode \t\t:= {flt}\t,\n"
            f"             oLft \t\t=> {o_lft}\t,\n"
            f"             oPrsLo \t\t=> {o_lo}\t,\n"
            f"             oBrakePrs \t\t=> {prs}\t,\n"
            f"\t     irst \t\t:= {rst});\n"
        )
        for z in b.zs_lifts:
            zt = sim_tag(z, "_I")
            need(zt, "EBOOL")
            st.append(f"{zt} := {o_lft};")
        st.append("")
    if not brakes:
        st.append("(* No BK units with SV+PS found *)\n")

    # ---- CAX from SR ----
    st.append(banner("Contactor aux (CAX) follows start relay (SR)"))
    stats.categories["CAX_from_SR"] = len(cax_sr)
    for cax_b, sr_b in cax_sr:
        ct = sim_tag(cax_b, "_I")
        so = do_tag(sr_b)
        need(ct, "EBOOL")
        need(so, "EBOOL")
        st.append(f"{ct} \t:= {so};")
    st.append("")
    if not cax_sr:
        st.append("(* No CAX+SR pairs *)\n")

    # ---- CB Simulation ----
    st.append(banner("CB Simulation"))
    stats.categories["sim_CBSts"] = len(cb_pairs)
    for i, p in enumerate(cb_pairs):
        st.append(f"(* {p.label} *)")
        set_trp = f"sim_{p.trp_base}_SetTrp"
        set_ok = f"sim_{p.trp_base}_SetOK"
        set_opn = f"sim_{p.trp_base}_SetOpn"
        cb_i = sim_tag(p.cb_base, "_I")
        trp_i = sim_tag(p.trp_base, "_I")
        need(set_trp, "EBOOL", f"{p.trp_base} sim set tripped")
        need(set_ok, "EBOOL", f"{p.trp_base} sim set CB closed/healthy")
        need(set_opn, "EBOOL", f"{p.trp_base} sim set CB open")
        need(cb_i, "EBOOL")
        need(trp_i, "EBOOL")
        need(f"sim_CBSts_{i}", "sim_CBSts")
        st.append(
            f"sim_CBSts_{i} (ioSetTripped := {set_trp},\n"
            f"             ioReset \t:= {set_ok},\n"
            f"             ioSetOpen \t:= {set_opn},\n"
            f"             oClsd \t=> {cb_i},\n"
            f"             oTrpd \t=> {trp_i});\n"
        )
    if not cb_pairs:
        st.append("(* No CB+TRP pairs found *)\n")

    st.append("")
    st.append("END_IF;\n")
    st.append("")
    st.append(banner("Equipment Simulations - END"))

    # dataBlock
    db_lines: list[str] = []
    for name in sorted(db_names):
        typ, comment = db_names[name]
        if typ == "EBOOL":
            db_lines.append(var_ebool(name, comment))
        else:
            db_lines.append(var_typed(name, typ, comment))

    st_text = "\n".join(st)
    # tidy double blanks
    st_text = re.sub(r"\n{3,}", "\n\n", st_text)
    return st_text, "".join(db_lines), stats



# ---------------------------------------------------------------------------
# IED / PR-rack tags for sim_Initialise (skipped by X80-only IOmap)
# ---------------------------------------------------------------------------

def collect_ied_init_tags(site_name: str) -> dict[str, list[tuple[str, int, str]]]:
    """Collect sim_ DI tags for TeSysT + REX640 (+ ensure FlexiSoft field DI on BSR132).

    Returns category → list of (tag, init_val, comment).
    """
    out: dict[str, list[tuple[str, int, str]]] = {
        "TeSysT": [],
        "REX640": [],
        "FlexiSoft": [],
    }
    seen: set[str] = set()

    def add(cat: str, tag: str, val: int, comment: str):
        if not tag or tag in seen:
            return
        seen.add(tag)
        out[cat].append((tag, val, comment))

    for u in collect_tesyst(site_name):
        if u.cb_base:
            add("TeSysT", sim_tag(u.cb_base, "_I"), u.cb_init, f"{u.drive} TeSysT CB01")
        if u.cax_base:
            add("TeSysT", sim_tag(u.cax_base, "_I"), u.cax_init, f"{u.drive} TeSysT CAX01")
        if u.sy_base:
            add("TeSysT", sim_tag(u.sy_base, "_I"), u.sy_init, f"{u.drive} TeSysT SY01")
        if u.vac_base:
            add("TeSysT", sim_tag(u.vac_base, "_I"), u.vac_init, f"{u.drive} TeSysT VAC01")

    if site_name == "BSR132":
        for rel in collect_rex640(site_name):
            for c in rel.channels:
                if c.iot not in ("DI", "DIO") or not c.device:
                    continue
                base = rex_sim_base(c.drive, c.device, "")
                if not base:
                    continue
                tag = sim_tag(base, "_I")
                val = state_to_val(c.state, c.device, c.desc)
                add("REX640", tag, val, f"{rel.owner} Slot {c.slot} {c.ch} {c.device}")

        # FlexiSoft / PNOZ-related field DI already discovered for safety — ensure present
        _flexi, _pnoz, feedbacks = collect_flexisoft_and_safety(site_name)
        for fb in feedbacks:
            tag = fb.sim_i or (sim_tag(fb.base, "_I") if fb.base else None)
            if tag and str(tag).startswith("sim_"):
                add("FlexiSoft", tag, 1, f"{fb.drive} {fb.device} safety DI".strip())

    return out


def _extract_xst_st(xst: str) -> str:
    m = re.search(r"<STSource>(.*?)</STSource>", xst, re.S)
    if not m:
        return ""
    st = m.group(1)
    return st.replace("&gt;", ">").replace("&lt;", "<").replace("&amp;", "&")


def _replace_xst_st_and_merge_vars(xst: str, new_st: str, new_var_xml: str) -> str:
    """Replace STSource and merge variable declarations into dataBlock."""
    st_xml = new_st.replace("&", "&amp;").replace("=>", "=&gt;")
    xst = re.sub(
        r"<STSource>.*?</STSource>",
        f"<STSource>{st_xml}</STSource>",
        xst,
        count=1,
        flags=re.S,
    )
    # Collect existing var names
    existing = {m.group(1).lower() for m in re.finditer(r'<variables name="([^"]+)"', xst)}
    to_add = []
    for m in re.finditer(
        r'(<variables name="([^"]+)"[^>]*(?:/>|>.*?</variables>))',
        new_var_xml,
        re.S,
    ):
        block, name = m.group(1), m.group(2)
        if name.lower() not in existing:
            # normalize to tab-indented block ending with newline
            if not block.endswith("\n"):
                block = block + "\n"
            if not block.startswith("\t\t"):
                block = "\t\t" + block.lstrip()
            to_add.append(block if block.startswith("\t\t<variables") else f"\t\t{block.lstrip()}")
            existing.add(name.lower())
    if to_add:
        # Insert before </dataBlock>
        insert = "".join(to_add)
        xst = xst.replace("\t</dataBlock>", insert + "\t</dataBlock>", 1)
    return xst


def refresh_initialise_with_ied(site_name: str) -> dict:
    """Regenerate sim_Initialise (X80 DI + IED/PR TeSysT/REX via generate_xst hook)."""
    from generate_xst import generate_initialise

    site = load_site(site_name)
    ied = collect_ied_init_tags(site_name)
    init_xml, init_stats = generate_initialise(site)

    out_path = OUTPUTS / site_name / "sim_Initialise.XST"
    out_path.write_text(init_xml, encoding="utf-8")

    # Count how many IED tags landed in the written ST
    st = _extract_xst_st(init_xml)
    present = set(re.findall(r"^(sim_\S+)\s*:=", st, re.M))
    counts = {cat: len(items) for cat, items in ied.items()}
    new_counts = {
        cat: sum(1 for t, _v, _c in items if t in present)
        for cat, items in ied.items()
    }
    return {
        "path": str(out_path),
        "bytes": out_path.stat().st_size,
        "ied_counts": counts,
        "ied_new_counts": new_counts,
        "total_init_tags": init_stats.get("init_tags"),
        "ied_tags_in_init_stats": init_stats.get("ied_tags"),
        "base_init": init_stats,
    }


TESYST_SECTION_RE = re.compile(
    r"\(\* =+\s*\n\s*\n\s*TeSysT motor starters \(SIM_TeSysT\).*?"
    r"(?=\(\* =+\s*\n\s*\n\s*(?:Equipment Simulations - END|Reference patterns|Analog AI))",
    re.S,
)


def _tesyst_section_text(units: list[TeSysTUnit], need, stats: GenStats) -> str:
    lines = emit_tesyst_st(units, need, stats)
    return "\n".join(lines).rstrip() + "\n\n"


def write_site(site_name: str) -> dict:
    site = load_site(site_name)
    out_dir = OUTPUTS / site_name
    out_dir.mkdir(parents=True, exist_ok=True)
    xst_path = out_dir / "sim_Equipment.XST"
    st_path = out_dir / "sim_Equipment.ST"

    # BSR132: preserve Tim's non-TeSysT edits — only replace TeSysT section + merge vars
    if site_name == "BSR132" and xst_path.exists() and st_path.exists():
        stats = GenStats()
        db_names: dict[str, tuple[str, Optional[str]]] = {}

        def need(name: str, typ: str, comment: str | None = None):
            if name not in db_names:
                db_names[name] = (typ, comment)

        units = collect_tesyst(site_name)
        tesyst_st = _tesyst_section_text(units, need, stats)
        # Also collect REX for report completeness (section already in file)
        rex = collect_rex640(site_name)
        if rex:
            stats.categories["sim_rex640"] = len(rex)
            stats.rex_report = {"count": len(rex), "owners": [r.owner for r in rex]}

        old_st = st_path.read_text(encoding="utf-8")
        if TESYST_SECTION_RE.search(old_st):
            new_st = TESYST_SECTION_RE.sub(tesyst_st, old_st, count=1)
        else:
            # Insert before Equipment Simulations - END
            end_banner = re.compile(
                r"\(\* =+\s*\n\s*\n\s*Equipment Simulations - END",
                re.S,
            )
            m = end_banner.search(old_st)
            if not m:
                raise RuntimeError("BSR132 sim_Equipment.ST missing END banner")
            new_st = old_st[: m.start()] + tesyst_st + old_st[m.start() :]

        # Drop obsolete wrong-map SY decls are ok to leave; ensure new tags declared
        db_lines = []
        for name in sorted(db_names):
            typ, comment = db_names[name]
            if typ == "EBOOL":
                db_lines.append(var_ebool(name, comment))
            else:
                db_lines.append(var_typed(name, typ, comment))
        new_var_xml = "".join(db_lines)

        old_xst = xst_path.read_text(encoding="utf-8")
        # Keep program identity / headers; replace ST + merge vars
        new_xst = _replace_xst_st_and_merge_vars(old_xst, new_st, new_var_xml)

        # Remove hardcoded leftovers check is done in report
        st_path.write_text(new_st, encoding="utf-8")
        xst_path.write_text(new_xst, encoding="utf-8")
    else:
        st, db, stats = emit_equipment(site)
        xst = wrap_xst(site.bpl, "sim_Equipment", SECTION_ORDER, st, db)
        xst_path.write_text(xst, encoding="utf-8")
        st_path.write_text(st, encoding="utf-8")

    init_report = refresh_initialise_with_ied(site_name)
    stats.init_ied_report = init_report.get("ied_counts", {})

    return {
        "site": site_name,
        "bpl": site.bpl,
        "xst": str(xst_path),
        "st": str(st_path),
        "xst_bytes": xst_path.stat().st_size,
        "st_bytes": st_path.stat().st_size,
        "categories": stats.categories,
        "skipped": stats.skipped,
        "rex_report": stats.rex_report,
        "safety_report": stats.safety_report,
        "tesyst_report": stats.tesyst_report,
        "init_report": init_report,
    }


def main():
    reports = []
    for site in ("BSR130", "BSR132"):
        reports.append(write_site(site))
    for r in reports:
        print("=" * 60)
        print(f"{r['site']} ({r['bpl']})")
        print(f"  XST: {r['xst']} ({r['xst_bytes']} bytes)")
        print(f"  ST:  {r['st']} ({r['st_bytes']} bytes)")
        print(f"  categories: {r['categories']}")
        print("  skipped:")
        for s in r["skipped"]:
            print(f"    - {s}")
        rr = r.get("rex_report") or {}
        print(f"  REX640: count={rr.get('count')} owners={rr.get('owners')}")
        print(f"  REX640 bit mapping: {rr.get('bit_mapping')}")
        um = rr.get("unmapped") or []
        print(f"  REX640 unmapped/FALSE channels: {len(um)}")
        sr = r.get("safety_report") or {}
        print(f"  Safety FlexiSoft owners: {sr.get('flexi_owners')} (BSR132-only per Tim)")
        print(f"  Safety PNOZ types: {sr.get('pnoz_types')}")
        print(f"  Safety feedbacks: {sr.get('feedback_count')}")
        tr = r.get("tesyst_report") or {}
        print(f"  TeSysT: count={tr.get('count')} owners={tr.get('owners')}")
        print(f"  TeSysT CB/CAX/SY/VAC: {tr.get('cb')}/{tr.get('cax')}/{tr.get('sy')}/{tr.get('vac')}")
        print(f"  TeSysT wiring: {tr.get('wiring')}")
        ir = r.get("init_report") or {}
        print(f"  Initialise: {ir.get('path')} ({ir.get('bytes')} bytes)")
        print(f"  Initialise IED collected: {ir.get('ied_counts')}")
        print(f"  Initialise IED newly added: {ir.get('ied_new_counts')}")
        print(f"  Initialise total tags: {ir.get('total_init_tags')}")


if __name__ == "__main__":
    main()
